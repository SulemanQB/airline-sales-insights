"""Build the monthly commercial pack as an Excel workbook with openpyxl."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config as C

HDR = PatternFill("solid", fgColor="14365D")
HFONT = Font(color="FFFFFF", bold=True)
GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
MONEY, PCT, NUM = '"$"#,##0', "0.0%", "#,##0"


def _table(ws, df: pd.DataFrame, row: int, col: int = 1, fmts: dict | None = None) -> tuple[int, int]:
    fmts = fmts or {}
    for j, c in enumerate(df.columns):
        cell = ws.cell(row=row, column=col + j, value=str(c))
        cell.fill, cell.font = HDR, HFONT
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, rec in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(rec):
            if isinstance(v, pd.Timestamp):
                v = v.to_pydatetime().date()
            if pd.isna(v) if not isinstance(v, (list, dict)) else False:
                v = None
            cell = ws.cell(row=row + i, column=col + j, value=v)
            f = fmts.get(df.columns[j])
            if f:
                cell.number_format = f
    for j, c in enumerate(df.columns):
        ws.column_dimensions[get_column_letter(col + j)].width = max(12, len(str(c)) + 2)
    return row + len(df), col + len(df.columns) - 1


def build_pack(con, res: dict, charts: dict, out: Path, month: str | None = None) -> Path:
    k = res["kpi"].copy()
    k["month_start"] = pd.to_datetime(k.month_start)
    m = pd.Timestamp(month) if month else k.month_start.max()
    cur = k[k.month_start == m].iloc[0]
    prev = k[k.month_start == m - pd.offsets.MonthBegin(1)].iloc[0]
    ly = k[k.month_start == m - pd.DateOffset(years=1)]
    ly = ly.iloc[0] if len(ly) else None
    wb = Workbook()

    # ---- Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = f"{C.AIRLINE}: Sales & Distribution commercial pack, {m:%B %Y}"
    ws["A1"].font = Font(size=14, bold=True, color="14365D")
    ws["A2"] = "Synthetic data. All figures regenerate from `make report`."
    metrics = [("Gross revenue", "gross_revenue", MONEY), ("Net revenue (after refunds)", "net_revenue", MONEY),
               ("Distribution cost", "distribution_cost", MONEY), ("Net revenue after distribution", "net_after_dist", MONEY),
               ("Passengers", "pax", NUM), ("Yield (c/RPK)", "yield_cents_rpk", "0.00"),
               ("RASK proxy (c/ASK)", "rask_proxy_cents", "0.000"), ("Cost of sale %", "cost_of_sale_pct", "0.00%"),
               ("Target attainment", "target_attainment", PCT)]
    rows = []
    for label, col, _ in metrics:
        a, b = cur[col], prev[col]
        c = ly[col] if ly is not None else None
        rows.append([label, a, b, a / b - 1 if b else None, c, (a / c - 1) if c else None])
    df = pd.DataFrame(rows, columns=["Metric", "This month", "Prior month", "MoM %", "Same month LY", "YoY %"])
    end, _ = _table(ws, df, 4)
    for i, (_, _, f) in enumerate(metrics, start=5):
        for col in "BCE":
            ws[f"{col}{i}"].number_format = f
        for col in "DF":
            ws[f"{col}{i}"].number_format = PCT
    # green = good. For cost of sale a fall is good, so that row gets flipped colours.
    cos_row = 5 + [x[1] for x in metrics].index("cost_of_sale_pct")
    for col in "DF":
        for r in range(5, end + 1):
            good_if_up = r != cos_row
            ws.conditional_formatting.add(f"{col}{r}", CellIsRule(operator="lessThan", formula=["0"], fill=RED if good_if_up else GREEN))
            ws.conditional_formatting.add(f"{col}{r}", CellIsRule(operator="greaterThanOrEqual", formula=["0"], fill=GREEN if good_if_up else RED))
    ws.column_dimensions["A"].width = 32
    r0 = end + 2
    ws.cell(row=r0, column=1, value="Recommendations").font = Font(bold=True, size=12, color="14365D")
    for i, rec in enumerate(res["recs"], start=1):
        ws.cell(row=r0 + i, column=1, value=f"{i}. {rec['title']}").font = Font(bold=True)
        ws.cell(row=r0 + i, column=2, value=rec["impact_aud"]).number_format = MONEY
        ws.cell(row=r0 + i, column=3, value=rec["impact_text"])
    img = XLImage(str(charts["revenue_vs_target"])); img.width, img.height = 640, 275
    ws.add_image(img, "H4")

    # ---- Trend (all months) with native line chart
    wt = wb.create_sheet("Trend")
    cols = ["month_start", "gross_revenue", "target_revenue", "target_attainment", "net_after_dist",
            "cost_of_sale_pct", "yield_cents_rpk", "rask_proxy_cents", "revenue_yoy_pct", "revenue_mom_pct"]
    tdf = k[cols].copy(); tdf["month_start"] = tdf.month_start.dt.date
    end, last_col = _table(wt, tdf, 1, fmts={"gross_revenue": MONEY, "target_revenue": MONEY, "net_after_dist": MONEY,
                                             "target_attainment": PCT, "cost_of_sale_pct": "0.00%", "yield_cents_rpk": "0.00",
                                             "rask_proxy_cents": "0.000", "revenue_yoy_pct": PCT, "revenue_mom_pct": PCT})
    for c in "A":
        for r in range(2, end + 1):
            wt[f"{c}{r}"].number_format = "mmm-yy"
    wt.conditional_formatting.add(f"D2:D{end}", ColorScaleRule(start_type="num", start_value=0.85, start_color="F8696B",
                                                              mid_type="num", mid_value=1.0, mid_color="FFEB84",
                                                              end_type="num", end_value=1.05, end_color="63BE7B"))
    lc = LineChart(); lc.title = "Gross revenue vs target"; lc.height, lc.width = 8, 18
    lc.add_data(Reference(wt, min_col=2, max_col=3, min_row=1, max_row=end), titles_from_data=True)
    lc.set_categories(Reference(wt, min_col=1, min_row=2, max_row=end))
    wt.add_chart(lc, "L2")

    # ---- Channel pivot (last 13 months)
    wc = wb.create_sheet("Channel Mix")
    cm = con.execute("SELECT month_start, channel || ' (' || connection || ')' AS ch, gross_revenue, revenue_share, "
                     "cost_of_sale_pct FROM v_channel_mix_monthly").df()
    cm["month_start"] = pd.to_datetime(cm.month_start)
    cm = cm[(cm.month_start > m - pd.DateOffset(months=13)) & (cm.month_start <= m)]
    piv = cm.pivot_table(index="month_start", columns="ch", values="gross_revenue", aggfunc="sum").reset_index()
    piv["month_start"] = piv.month_start.dt.strftime("%b-%y")
    piv.columns.name = None
    end, last_col = _table(wc, piv.rename(columns={"month_start": "Month"}), 1,
                           fmts={c: MONEY for c in piv.columns if c != "month_start"})
    bc = BarChart(); bc.type = "col"; bc.grouping = "percentStacked"; bc.overlap = 100
    bc.title = "Revenue share by channel"; bc.height, bc.width = 9, 20
    bc.add_data(Reference(wc, min_col=2, max_col=last_col, min_row=1, max_row=end), titles_from_data=True)
    bc.set_categories(Reference(wc, min_col=1, min_row=2, max_row=end))
    wc.add_chart(bc, f"A{end + 3}")
    sh = cm.pivot_table(index="month_start", columns="ch", values="cost_of_sale_pct").reset_index()
    sh["month_start"] = sh.month_start.dt.strftime("%b-%y"); sh.columns.name = None
    r2 = end + 23
    wc.cell(row=r2 - 1, column=1, value="Cost of sale % by channel").font = Font(bold=True)
    e2, lc2 = _table(wc, sh.rename(columns={"month_start": "Month"}), r2, fmts={c: "0.0%" for c in sh.columns})
    wc.conditional_formatting.add(f"B{r2 + 1}:{get_column_letter(lc2)}{e2}",
                                  ColorScaleRule(start_type="min", start_color="63BE7B", end_type="max", end_color="F8696B"))

    # ---- Routes for the month
    wr = wb.create_sheet("Routes")
    rt = con.execute("SELECT * FROM v_route_monthly").df()
    rt["month_start"] = pd.to_datetime(rt.month_start)
    rt = rt[rt.month_start == m][["route_code", "haul", "pax", "gross_revenue", "target_revenue", "target_attainment",
                                  "revenue_yoy_pct", "net_after_dist", "yield_cents_rpk", "rask_proxy_cents"]]
    rt = rt.sort_values("target_attainment")
    rt["variance_to_target"] = rt.gross_revenue - rt.target_revenue
    end, _ = _table(wr, rt, 1, fmts={"pax": NUM, "gross_revenue": MONEY, "target_revenue": MONEY, "net_after_dist": MONEY,
                                     "target_attainment": PCT, "revenue_yoy_pct": PCT, "yield_cents_rpk": "0.00",
                                     "rask_proxy_cents": "0.000", "variance_to_target": MONEY})
    wr.conditional_formatting.add(f"F2:F{end}", CellIsRule(operator="lessThan", formula=["0.9"], fill=RED))
    wr.conditional_formatting.add(f"F2:F{end}", CellIsRule(operator="greaterThanOrEqual", formula=["1"], fill=GREEN))
    wr.conditional_formatting.add(f"K2:K{end}", DataBarRule(start_type="min", end_type="max", color="1B9AAA"))
    b2 = BarChart(); b2.type = "bar"; b2.title = "Target attainment by route"; b2.height, b2.width = 9, 16
    b2.add_data(Reference(wr, min_col=6, min_row=1, max_row=end), titles_from_data=True)
    b2.set_categories(Reference(wr, min_col=1, min_row=2, max_row=end))
    wr.add_chart(b2, "M2")

    # ---- Lead time cohort pivot
    wl = wb.create_sheet("Lead Time")
    lt = con.execute("SELECT * FROM v_leadtime_cohort").df()
    lt["month_start"] = pd.to_datetime(lt.month_start)
    lt = lt[lt.month_start == m]
    p1 = lt.pivot_table(index="lead_bucket", columns="segment", values="bookings", aggfunc="sum").reset_index()
    p2 = lt.pivot_table(index="lead_bucket", columns="segment", values="avg_fare_per_pax").reset_index()
    order = ["0-7", "8-21", "22-60", "61+"]
    for p in (p1, p2):
        p.columns.name = None
        p["lead_bucket"] = pd.Categorical(p.lead_bucket, order)
        p.sort_values("lead_bucket", inplace=True)
        p["lead_bucket"] = p.lead_bucket.astype(str)
    wl["A1"] = "Bookings by lead time (days before travel)"; wl["A1"].font = Font(bold=True)
    e1, _ = _table(wl, p1.rename(columns={"lead_bucket": "Lead days"}), 2, fmts={"Corporate": NUM, "Leisure": NUM})
    wl.cell(row=e1 + 2, column=1, value="Average fare per passenger").font = Font(bold=True)
    e2, _ = _table(wl, p2.rename(columns={"lead_bucket": "Lead days"}), e1 + 3, fmts={"Corporate": MONEY, "Leisure": MONEY})
    wl.conditional_formatting.add(f"B{e1 + 4}:C{e2}", ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="1B9AAA"))

    # ---- Anomalies + forecast
    wa = wb.create_sheet("Alerts & Forecast")
    fl = res["anomaly"]["flagged"].copy()
    fl["booking_date"] = fl.booking_date.dt.date
    wa["A1"] = "Daily channel anomalies (bookings far below same-weekday norm)"; wa["A1"].font = Font(bold=True)
    e1, _ = _table(wa, fl.round(2), 2, fmts={"expected": "0.0", "z": "0.0"})
    fc = res["forecast"]["forecast"].copy(); fc["month_start"] = fc.month_start.dt.date
    wa.cell(row=e1 + 2, column=1, value=f"Next-quarter forecast ({res['forecast']['best_model']}, backtest MAPE "
                                        f"{res['forecast']['models'][res['forecast']['best_model']]:.1f}%)").font = Font(bold=True)
    e2, _ = _table(wa, fc, e1 + 3, fmts={"gross_revenue": MONEY})
    for r in range(e1 + 4, e2 + 1):
        wa[f"A{r}"].number_format = "mmm-yy"
    img = XLImage(str(charts["forecast"])); img.width, img.height = 600, 257
    wa.add_image(img, "H2")

    # ---- Definitions
    wd = wb.create_sheet("Definitions")
    defs = [("Gross revenue", "Fare + ancillary revenue at time of booking"),
            ("Net revenue", "Gross revenue minus refunds"),
            ("Distribution cost", "Agency/OTA commission + GDS or NDC per-segment fee + card fees on direct sales"),
            ("Cost of sale %", "Distribution cost / gross revenue"),
            ("Yield", "Fare revenue per revenue passenger kilometre (cents)"),
            ("RASK proxy", "Net revenue per available seat kilometre (cents); seats are a fixed daily schedule"),
            ("Target attainment", "Gross revenue / monthly plan target")]
    _table(wd, pd.DataFrame(defs, columns=["Term", "Definition"]), 1)
    wd.column_dimensions["B"].width = 90
    for w in wb.worksheets:
        w.freeze_panes = None
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out
